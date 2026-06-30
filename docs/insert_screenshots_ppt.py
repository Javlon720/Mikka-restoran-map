from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

ppt_path = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/cowork/Mikka_Presentation.pptx"
img_dir = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/cowork"

def main():
    if not os.path.exists(ppt_path):
        print("Mikka_Presentation.pptx not found.")
        return

    prs = Presentation(ppt_path)
    
    # Add new slide for screenshots
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background style
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 247, 250)
    
    # Slide Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "デモンストレーション（システム稼働画面）"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.name = 'MS Gothic'
    p.font.color.rgb = RGBColor(30, 41, 59)
    
    # Images config
    # 2x2 grid layout coordinates on standard widescreen/4:3 slide (we assume safe margins)
    images = [
        ("screenshot_main.png", "メイン画面（地図とレストラン一覧）", Inches(0.5), Inches(1.1), Inches(4.2), Inches(2.5)),
        ("screenshot_chat.png", "Mikka AI Assistant（Geminiチャット）", Inches(5.3), Inches(1.1), Inches(4.2), Inches(2.5)),
        ("screenshot_login.png", "多言語ログイン・登録フォーム", Inches(0.5), Inches(4.3), Inches(4.2), Inches(2.5)),
        ("screenshot_add.png", "飲食店追加・位置指定（地図連携）", Inches(5.3), Inches(4.3), Inches(4.2), Inches(2.5))
    ]
    
    for filename, caption, left, top, width, height in images:
        img_path = os.path.join(img_dir, filename)
        if os.path.exists(img_path):
            # Insert picture
            slide.shapes.add_picture(img_path, left, top, width=width, height=height)
            
            # Insert caption textbox below picture
            capBox = slide.shapes.add_textbox(left, top + height, width, Inches(0.4))
            cap_tf = capBox.text_frame
            cap_tf.word_wrap = True
            cap_tf.margin_top = Inches(0.02)
            cap_p = cap_tf.paragraphs[0]
            cap_p.text = caption
            cap_p.font.size = Pt(11)
            cap_p.font.name = 'MS Gothic'
            cap_p.font.color.rgb = RGBColor(71, 85, 105)
            cap_p.alignment = PP_ALIGN.CENTER
            
    prs.save(ppt_path)
    print("Successfully inserted screenshots into Mikka_Presentation.pptx.")

if __name__ == "__main__":
    main()
