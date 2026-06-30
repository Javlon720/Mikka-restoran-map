from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1: Title
blank_slide_layout = prs.slide_layouts[6]
slide1 = prs.slides.add_slide(blank_slide_layout)

# Background color or simple clean style
background = slide1.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(245, 247, 250)

# Add Title Text
txBox = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "自己紹介"
p.font.size = Pt(44)
p.font.bold = True
p.font.name = 'MS Gothic'
p.font.color.rgb = RGBColor(30, 41, 59)

p2 = tf.add_paragraph()
p2.text = "アブドゥラエフ・ジャブロン (Abdullayev Javlon)\nJapan Digital University (JDU) / コンピュータ工学専攻"
p2.font.size = Pt(20)
p2.font.name = 'MS Gothic'
p2.font.color.rgb = RGBColor(100, 116, 139)

# Slide 2: Strength 1
slide2 = prs.slides.add_slide(blank_slide_layout)
slide2.background.fill.solid()
slide2.background.fill.fore_color.rgb = RGBColor(245, 247, 250)

txBox = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "強み①：フルスタックなWeb開発力"
p.font.size = Pt(32)
p.font.bold = True
p.font.name = 'MS Gothic'
p.font.color.rgb = RGBColor(30, 41, 59)

bullets = [
    "• 主要スキル: Node.js, TypeScript, JavaScript, PostgreSQL, React",
    "• 設計思想の重視: クリーンアーキテクチャの導入による高保守性システムの構築",
    "• コワーク開発実績: 飲食店情報口コミサイト『Mikka』の開発",
    "   - Leaflet.js地図APIとのシームレスな位置情報連携",
    "   - Google Gemini APIを用いたAIチャットボットによるレコメンド機能",
    "   - 多言語対応（日本語・ウズベク語・英語・ロシア語）の実装",
    "   - Google OAuth 2.0認証とJWTセッション管理による高いセキュリティ"
]
for bullet in bullets:
    p = tf.add_paragraph()
    p.text = bullet
    p.font.size = Pt(16)
    p.font.name = 'MS Gothic'
    p.font.color.rgb = RGBColor(71, 85, 105)

# Slide 3: Strength 2
slide3 = prs.slides.add_slide(blank_slide_layout)
slide3.background.fill.solid()
slide3.background.fill.fore_color.rgb = RGBColor(245, 247, 250)

txBox = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "強み②：日本語能力 ✕ デザイン表現力"
p.font.size = Pt(32)
p.font.bold = True
p.font.name = 'MS Gothic'
p.font.color.rgb = RGBColor(30, 41, 59)

bullets3 = [
    "• 高度な日本語能力: JLPT N2合格。日本人のパートナーと直接ビジネス交渉可能",
    "• 実務経験（グラフィックデザイン）:",
    "   - JDU公式ソーシャルメディア用ビジュアル作成",
    "   - 日本向け輸出Tシャツのデザインを1,000件以上担当した実績",
    "   - 新規事業のブランドブックやロゴデザインの構築支援",
    "• 教育・コミュニケーション能力:",
    "   - 日本語教師として20名以上の学生（子供から大人）にJLPT N5レベルを指導"
]
for bullet in bullets3:
    p = tf.add_paragraph()
    p.text = bullet
    p.font.size = Pt(16)
    p.font.name = 'MS Gothic'
    p.font.color.rgb = RGBColor(71, 85, 105)

# Slide 4: Future Goal
slide4 = prs.slides.add_slide(blank_slide_layout)
slide4.background.fill.solid()
slide4.background.fill.fore_color.rgb = RGBColor(245, 247, 250)

txBox = slide4.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "将来の目標"
p.font.size = Pt(32)
p.font.bold = True
p.font.name = 'MS Gothic'
p.font.color.rgb = RGBColor(30, 41, 59)

bullets4 = [
    "• 日本とウズベキスタンの企業を繋ぐ『ブリッジシステムエンジニア』としての活躍",
    "• 技術的な開発だけでなく、多言語コミュニケーションとUI/UXデザインを活かした",
    "  グローバルプロダクトの創出",
    "• 常に最新技術（AI連携、クリーンコード等）を学習し、組織の成長に貢献する"
]
for bullet in bullets4:
    p = tf.add_paragraph()
    p.text = bullet
    p.font.size = Pt(18)
    p.font.name = 'MS Gothic'
    p.font.color.rgb = RGBColor(71, 85, 105)

prs.save('/Users/macbook-projavlon/Desktop/Jdu/docs/7_自己紹介プレゼン資料.pptx')
print("Successfully generated PPT presentation.")
