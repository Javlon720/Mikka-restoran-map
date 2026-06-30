import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import copy

def copy_cell_style(src_cell, dst_cell):
    """Copies font, fill, border, and alignment from src_cell to dst_cell."""
    if src_cell.font:
        dst_cell.font = Font(
            name=src_cell.font.name,
            size=src_cell.font.size,
            bold=src_cell.font.bold,
            italic=src_cell.font.italic,
            color=copy.copy(src_cell.font.color),
            underline=src_cell.font.underline
        )
    if src_cell.fill:
        dst_cell.fill = PatternFill(
            fill_type=src_cell.fill.fill_type,
            start_color=copy.copy(src_cell.fill.start_color),
            end_color=copy.copy(src_cell.fill.end_color)
        )
    if src_cell.border:
        dst_cell.border = Border(
            left=copy.copy(src_cell.border.left),
            right=copy.copy(src_cell.border.right),
            top=copy.copy(src_cell.border.top),
            bottom=copy.copy(src_cell.border.bottom)
        )
    if src_cell.alignment:
        dst_cell.alignment = Alignment(
            horizontal=src_cell.alignment.horizontal,
            vertical=src_cell.alignment.vertical,
            wrap_text=src_cell.alignment.wrap_text,
            shrink_to_fit=src_cell.alignment.shrink_to_fit,
            indent=src_cell.alignment.indent
        )

# Load workbook
excel_path = '/Users/macbook-projavlon/Downloads/214773 - Abdullayev _Javlonbek.xlsx'
wb = openpyxl.load_workbook(excel_path)

# ==================== 1. 履歴書の更新 ====================
sheet_rireki = wb['履歴書']

# 職歴の更新
# 元々の Row 23 の Column D に職歴を詳細化
sheet_rireki.cell(row=23, column=4, value='グラフィックデザイナー、日本語教師等として従事（詳細は職務経歴書参照）')
sheet_rireki.cell(row=24, column=4, value='現在に至る　　　　　　　　　　　　　　　　　　以上')

# ==================== 2. 職務経歴書の更新 ====================
sheet_keireki = wb['職務経歴書']

# 既存の行からスタイルをコピーするための見本セルを取得
# Row 21 (その他職務経歴のヘッダー見本)
# Row 22 (その他職務経歴の詳細見本)

# A11:E11 のマージとヘッダーテキストの書き込み
sheet_keireki.merge_cells('A11:E11')
sheet_keireki['A11'].value = '2026年6月～2026年6月 ／ コワーク授業・補講プロジェクト（飲食店情報口コミサイト開発）'

# A12:C12 のマージと詳細テキストの書き込み
sheet_keireki.merge_cells('A12:C12')
sheet_keireki['A12'].value = (
    '【システム名】 飲食店情報口コミサイト「Mikka」の開発\n'
    '【作業内容】\n'
    '・多言語対応（日本語・ウズベク語・英語・ロシア語）を動的切り替え可能な自作i18nライブラリを実装。\n'
    '・Leaflet.js地図APIを連携し、店舗の位置情報を地図上にマッピングする機能およびドラッグ＆ドロップによる緯度・経度選択機能を実装。\n'
    '・Google Gemini APIを用いて、登録店舗データをコンテキストとして適応的におすすめ店舗を推薦するAIチャットボット「Mikka AI Assistant」を構築。\n'
    '・バックエンドにClean Architecture（Domain/Use Cases/Adapters/Infrastructure）を採用し、Node.js + TypeScript + PostgreSQLで拡張性の高いAPIサーバーを設計・実装。\n'
    '・Google OAuth 2.0によるソーシャルログイン認証およびJWTを用いたセッション管理（HTTP-only Cookie）を実装しセキュリティを確保。'
)

# D12: 使用技術
sheet_keireki['D12'].value = (
    '【使用技術・ツール】\n'
    'HTML5, CSS3, JavaScript, Vite,\n'
    'Node.js, Fastify, TypeScript,\n'
    'PostgreSQL, Gemini API,\n'
    'Leaflet.js, GitHub'
)

# E12: 役割
sheet_keireki['E12'].value = (
    '役割\n'
    'フルスタックデベロッパー\n'
    '（一人で全開発を担当）'
)

# スタイル（フォント、罫線、揃え方）のコピー適用
# Row 11 は Row 21 からコピー
for col in range(1, 6):
    src_c = sheet_keireki.cell(row=21, column=col)
    dst_c = sheet_keireki.cell(row=11, column=col)
    copy_cell_style(src_c, dst_c)
    # 太字に調整
    if dst_c.font:
        dst_c.font = Font(name=dst_c.font.name, size=dst_c.font.size, bold=True)

# Row 12 は Row 22 からコピー
for col in range(1, 6):
    src_c = sheet_keireki.cell(row=22, column=col)
    dst_c = sheet_keireki.cell(row=12, column=col)
    copy_cell_style(src_c, dst_c)
    # 折り返しとアライメント設定
    dst_c.alignment = Alignment(wrap_text=True, vertical='top')

# 行の高さを設定
sheet_keireki.row_dimensions[11].height = sheet_keireki.row_dimensions[21].height or 24
# 詳細部分はテキスト量が多いので高さを大きめにする
sheet_keireki.row_dimensions[12].height = 180

# 保存
wb.save(excel_path)
print("Excel update completed successfully!")
