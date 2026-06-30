from pptx import Presentation
import os

ppt_path = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/employability/abdullayevJavlon214773.pptx"

def main():
    if not os.path.exists(ppt_path):
        print(f"{ppt_path} not found.")
        return

    prs = Presentation(ppt_path)
    
    # Update Slide 1 text content (Graduated from World Languages University)
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            for p in tf.paragraphs:
                if "併学大学" in p.text:
                    p.text = "既卒大学: ウズベキスタン世界言語大学 コンピューター言語学専攻 (卒業)"
                    p.font.name = 'MS Gothic'
                    print("Updated Slide 1 text line for World Languages University graduation.")

    # Update Slide 1 speaker notes (N3 ます形)
    notes_slide1 = slide1.notes_slide
    notes_tf = notes_slide1.notes_text_frame
    original_notes = notes_tf.text
    # Replace simultaneous study with graduation
    updated_notes = original_notes.replace(
        "また、ウズベキスタン世界言語大学でコンピューター言語学も同時に勉強しています。",
        "また、ウズベキスタン世界言語大学のコンピューター言語学専攻をすでに卒業しております。"
    )
    notes_tf.text = updated_notes
    print("Updated Slide 1 speaker notes for graduation.")
    
    # Also update Slide 5 notes if it has similar reference
    slide5 = prs.slides[4]
    notes_slide5 = slide5.notes_slide
    notes_tf5 = notes_slide5.notes_text_frame
    original_notes5 = notes_tf5.text
    updated_notes5 = original_notes5.replace("同時に勉強しています", "卒業しております")
    notes_tf5.text = updated_notes5

    prs.save(ppt_path)
    print("Successfully updated graduation history in abdullayevJavlon214773.pptx!")

if __name__ == "__main__":
    main()
