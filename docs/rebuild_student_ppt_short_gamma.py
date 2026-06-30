from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

def create_card(slide, left, top, width, height, bg_color, border_color=None):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def add_clean_textbox(slide, left, top, width, height, text, font_size, font_color, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.name = 'MS Gothic'
    p.font.color.rgb = font_color
    p.alignment = align
    return tf

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors (Light Gamma Modern theme - Royal Blue & Teal Accent)
    bg_color = RGBColor(255, 255, 255)
    card_color = RGBColor(248, 250, 252)     # Slate-50
    border_color = RGBColor(226, 232, 240)   # Slate-200
    text_dark = RGBColor(15, 23, 42)
    text_gray = RGBColor(71, 85, 105)
    
    accent_blue = RGBColor(29, 78, 216)      # Blue-700
    accent_teal = RGBColor(13, 148, 136)     # Teal-600
    
    blank_layout = prs.slide_layouts[6]
    photo_path = '/Users/macbook-projavlon/Downloads/IMG_0844.JPG'
    has_photo = os.path.exists(photo_path)
    
    # 5枚のコンパクトで洗練されたスライドデータを定義
    slides_data = [
        # Slide 1: Title & Profile
        {
            "id": 1,
            "title": "自己紹介 (Self-Introduction)",
            "subtitle": "プロフィール",
            "content": [
                "氏名: ジャブロンベク・アブドゥラエフ・シュフラトオグリ",
                "年齢: 22歳 (2003年9月18日生まれ)",
                "所属: Japan Digital University (JDU) 5年生",
                "専攻: コンピューター工学科 (2026年卒業予定)",
                "併願: ウズベキスタン世界言語大学 (コンピューター言語学専攻)"
            ],
            "note": "【発表原稿（スライド1）】[N3・ます形] ※目安：30秒\n"
                    "はじめまして。アブドゥラエフ・ジャブロンと申します。本日はよろしくお願いいたします。\n"
                    "私は現在、ウズベキスタン世界言語大学でコンピューター言語学を勉強しています。同時に、Japan Digital University（JDU）でコンピューター工学を専攻しています。2026年に卒業する予定です。"
        },
        # Slide 2: What I study
        {
            "id": 2,
            "title": "大学での学習内容",
            "subtitle": "専門分野とスキル",
            "content": [
                "• プログラミング学習 (JavaScript, TypeScript, HTML/CSS)",
                "• データベースの基礎と設計 (PostgreSQL, SQLリレーション)",
                "• ネットワークの仕組みとセキュリティ (JWT認証, クッキー管理)",
                "• Web開発の基本とバックエンドシステム構築"
            ],
            "note": "【発表原稿（スライド2）】[N3・ます形] ※目安：45秒\n"
                    "大学で勉強していることについて説明します。\n"
                    "主にプログラミング、データベース設計、ネットワーク、そしてWeb開発について学びました。特にフロントエンドとバックエンドの両方を繋ぐWebの仕組みについて深く学習しています。"
        },
        # Slide 3: Strength 1 (Web Development)
        {
            "id": 3,
            "title": "強み①：Webシステム開発力",
            "subtitle": "コワーク活動での実績",
            "content": [
                "• 飲食店口コミサイト『Mikka（ミッカ）』を自分で開発",
                "  - Node.js, Fastify, TypeScript, PostgreSQLを使用",
                "  - Leaflet.jsを用いたインタラクティブな地図連携機能の実装",
                "  - Google Gemini APIを組み込んだAIチャットアシスタントの構築",
                "  - パスワードハッシュ化とJWTによる安全なユーザー認証の実装"
            ],
            "note": "【発表原稿（スライド3）】[N3・ます形] ※目安：60秒\n"
                    "私の1つ目の強みは、『自分でWebシステムを作る力』です。\n"
                    "最近のコワーク授業では、レストラン口コミサイト『Mikka（ミッカ）』を開発しました。このアプリは地図から店を探したり、GeminiのAIチャットでおすすめを聞いたりできます。TypeScriptやPostgreSQLを使い、セキュリティ設計まで一人で実装を行いました。"
        },
        # Slide 4: Strength 2 (Design & Language)
        {
            "id": 4,
            "title": "強み②：デザイン実績 ＆ 日本語能力",
            "subtitle": "多文化コミュニケーション能力",
            "content": [
                "• 日本向けTシャツのグラフィックデザイン実務経験",
                "  - 日本のパートナーと日本語でやり取りし、1,000件以上のデザインを制作",
                "• 日本語での高いコミュニケーション力",
                "  - 日本語能力試験（JLPT）N3合格 (N2を勉強中)",
                "• 日本語教師としての指導経験",
                "  - 他校で20名以上の学生に日本語(N5レベル)を指導"
            ],
            "note": "【発表原稿（スライド4）】[N3・ます形] ※目安：60秒\n"
                    "2つ目の強みは、『日本語能力とデザインの実務経験』です。\n"
                    "日本語はN3レベルで、JDUではグラフィックデザイナーとして働き、日本のパートナーと日本語で連絡を取りながら、日本向けのTシャツデザインを1,000件以上制作しました。また、日本語の先生として学生に教えた経験もあり、人と協力して仕事をすることが得意です。"
        },
        # Slide 5: Thank you
        {
            "id": 5,
            "title": "ご清聴ありがとうございました",
            "subtitle": "ご挨拶",
            "content": [
                "• お忙しい中、貴重なお時間をいただき本当にありがとうございました",
                "• これからも努力を続けて、成長していきます",
                "• 連絡先: abdullayevjavlon720@gmail.com",
                "• 氏名: ジャブロンベク・アブドゥラエフ"
            ],
            "note": "【発表原稿（スライド5）】[N3・ます形] ※目安：30秒\n"
                    "私は、これまでに身につけたWeb開発技術、デザインの実績、そして日本語能力を活かして、日本のチームで力になりたいと考えています。\n"
                    "本日はお忙しい中、貴重なお時間をいただき本当にありがとうございました。どうぞよろしくお願いいたします。"
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Background Pure White
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
        # 1. Slide Title
        add_clean_textbox(slide, Inches(1.0), Inches(0.5), Inches(11.3), Inches(0.8), data["title"], 28, accent_blue, bold=True)
        
        # Title Accent Line (Teal color)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.0), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_teal
        line.line.fill.background()
        
        # Layout (Slide 1 title has profile picture on the right)
        is_title_slide = data["id"] == 1
        
        if is_title_slide and has_photo:
            # 2 Columns: Left side Card for text, Right side for Photo
            create_card(slide, Inches(1.0), Inches(1.5), Inches(7.2), Inches(5.0), card_color, border_color)
            tf = add_clean_textbox(slide, Inches(1.4), Inches(1.8), Inches(6.4), Inches(4.4), data["subtitle"], 20, accent_teal, bold=True)
            
            for bullet in data["content"]:
                p = tf.add_paragraph()
                p.text = bullet
                p.font.size = Pt(14)
                p.font.name = 'MS Gothic'
                p.font.color.rgb = text_dark
                p.space_before = Pt(12)
                p.font.bold = not bullet.strip().startswith("•")
                
            # Photo (Right side with card container)
            photo_card = create_card(slide, Inches(8.6), Inches(1.5), Inches(3.7), Inches(5.0), card_color, border_color)
            slide.shapes.add_picture(photo_path, Inches(8.8), Inches(1.7), width=Inches(3.3), height=Inches(4.6))
        else:
            # Standard Slide: 1 Big Card
            create_card(slide, Inches(1.0), Inches(1.5), Inches(11.3), Inches(5.0), card_color, border_color)
            tf = add_clean_textbox(slide, Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.4), data["subtitle"], 20, accent_teal, bold=True)
            
            for bullet in data["content"]:
                p = tf.add_paragraph()
                p.text = bullet
                p.font.size = Pt(15)
                p.font.name = 'MS Gothic'
                p.font.color.rgb = text_dark
                p.space_before = Pt(12)
                p.font.bold = not bullet.strip().startswith("•") and not bullet.strip().startswith("-")
                
        # 5. Add Speaker Notes
        slide.notes_slide.notes_text_frame.text = data["note"]
        
    # Save as the correct filename requested by the user
    ppt_out = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/employability/abdullayevJavlon214773.pptx"
    prs.save(ppt_out)
    print("Rebuilt student resume PPTX into 5 compact and clean modern slides successfully!")

if __name__ == "__main__":
    main()
