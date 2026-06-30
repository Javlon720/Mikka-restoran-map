from pptx import Presentation
import os

ppt_path = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/employability/7_自己紹介プレゼン資料.pptx"

def main():
    if not os.path.exists(ppt_path):
        print("7_自己紹介プレゼン資料.pptx not found.")
        return

    prs = Presentation(ppt_path)
    
    self_intro_notes = [
        # Slide 1 (自己紹介タイトル)
        "【発表原稿（スライド1）】[N3レベル・ます形] ※制限時間：30秒〜45秒\n"
        "はじめまして。アブドゥラエフ・ジャブロンと申します。本日はよろしくお願いいたします。\n"
        "私は現在、ウズベキスタン世界言語大学でコンピュータ言語学を勉強しています。同時に、JDU（Japan Digital University）でコンピュータエンジニアリングを専攻しています。また、産業能率大学でビジネスも勉強しました。\n"
        "将来は日本やグローバル企業で、ソフトウェアエンジニアとして働きたいと考えています。\n"
        "本日は、私の強みを2点、お話しさせていただきます。",
        
        # Slide 2 (強み1：Web開発)
        "【発表原稿（スライド2）】[N3レベル・ます形] ※制限時間：60秒\n"
        "私の強みの1つ目は、『自分でWebシステムを作る力』です。\n"
        "私はJavaScriptやNode.js、PostgreSQLを使った開発が得意です。\n"
        "最近のコワーク授業では、レストランの口コミサイト『Mikka（ミッカ）』を自分で作りました。\n"
        "このサイトでは、地図で行きたい店を探したり、GeminiのAIチャットを使っておすすめの店を聞いたりすることができます。フロントエンドからバックエンドまで、一人で工夫して開発を行いました。",
        
        # Slide 3 (強み2：日本語・デザイン)
        "【発表原稿（スライド3）】[N3レベル・ます形] ※制限時間：60秒\n"
        "強みの2つ目は、『日本語でのコミュニケーション力とデザインの経験』です。\n"
        "私の日本語レベルはN3で、現在はN2の勉強をしています。日本語で話すことが大好きです。\n"
        "これまでJDUでグラフィックデザイナーとして働き、日本のパートナーと日本語で連絡を取りながら、日本向けのTシャツデザインを1,000件以上作りました。また、日本語の先生として、20人以上の生徒に日本語を教えた経験もあります。",
        
        # Slide 4 (目標)
        "【発表原稿（スライド4）】[N3レベル・ます形] ※制限時間：30秒〜45秒\n"
        "最後に、私の将来の目標についてです。\n"
        "将来は、この『Web開発の技術』『デザインの経験』、そして『日本語力』を活かして、日本とウズベキスタンを繋ぐブリッジシステムエンジニアとして、人の役に立つシステムを作りたいです。\n"
        "本日は私のことを知っていただけるよう、一生懸命話します。ありがとうございました。どうぞよろしくお願いいたします。"
    ]
    
    for i, slide in enumerate(prs.slides):
        if i < len(self_intro_notes):
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = self_intro_notes[i]
            print(f"Updated N3 self-intro notes for Slide {i+1}")
            
    prs.save(ppt_path)
    print("Successfully updated presenter notes for self-introduction slides in N3 style!")

if __name__ == "__main__":
    main()
