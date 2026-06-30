from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

ppt_path = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/cowork/Mikka_Presentation.pptx"
img_dir = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/cowork"

def main():
    if not os.path.exists(ppt_path):
        print("Mikka_Presentation.pptx not found.")
        return

    prs = Presentation(ppt_path)
    
    # 1. 先ほど追加した 2x2 の大雑把なデモスライドを削除する（もし存在すれば）
    # 通常、編集前の元スライド数は 5枚または 6枚程度であると想定されます。
    # 最後のスライドが「デモンストレーション」になっている場合は削除します。
    if len(prs.slides) > 0:
        last_slide = prs.slides[-1]
        # タイトルを確認して「デモンストレーション」なら削除
        has_demo_title = False
        for shape in last_slide.shapes:
            if shape.has_text_frame and "デモンストレーション" in shape.text:
                has_demo_title = True
                break
        if has_demo_title:
            # スライドを削除
            id_list = prs.slides._sldIdLst
            del id_list[-1]
            print("Removed the previous simple demo slide.")

    # 2. 詳細な説明スライドを 4枚追加する
    # 各スライドの共通背景色 (F5F7FA)
    blank_layout = prs.slide_layouts[6]
    
    slides_data = [
        {
            "title": "機能紹介①：メイン画面 ＆ マップ機能",
            "img": "screenshot_main.png",
            "bullets": [
                "【特徴】インタラクティブな地図連携と店舗検索",
                "・Leaflet.js地図APIとのシームレスな連携",
                "  データベースから取得した飲食店をリアルタイムで地図上にマッピングします。",
                "・地図とリストの連動アクション",
                "  店舗カードをクリックすると地図上の対象ピンがポップアップし、",
                "  詳細情報モーダルへ直感的に遷移できます。",
                "・リアルタイム検索・フィルタリング",
                "  店舗名や料理ジャンル（Uzbek, Japanese等）によるリアルタイム検索。"
            ]
        },
        {
            "title": "機能紹介②：多言語対応(i18n) ＆ ユーザー認証",
            "img": "screenshot_login.png",
            "bullets": [
                "【特徴】グローバルな利用環境と安全な認証",
                "・自作の動的多言語（i18n）ライブラリ",
                "  日本語・ウズベク語・英語・ロシア語の4言語を、",
                "  リロードなしで即時に切り替えることが可能です。",
                "・セキュアなユーザー認証機能",
                "  JWT（JSON Web Token）およびHTTP-only Cookieを用いたセッション管理。",
                "  パスワードはbcryptjsで高度に暗号化され保存されます。",
                "・Google OAuth 2.0 ソーシャルログイン対応",
                "  Googleアカウントによるワンクリックログイン機能を統合。"
            ]
        },
        {
            "title": "機能紹介③：管理者機能（店舗情報管理）",
            "img": "screenshot_add.png",
            "bullets": [
                "【特徴】管理者専用のコンテンツ追加・編集画面",
                "・ロールベースのアクセス制御（RBAC）",
                "  管理者（admin）のみに店舗の追加・編集・削除権限を付与。",
                "・直感的な位置指定マッピング",
                "  管理者画面内のマップでピンをドラッグ＆ドロップするだけで、",
                "  自動的に緯度・経度の座標が取得されフォームに入力されます。",
                "・Multerによる画像アップロード",
                "  店舗画像をアップロードし、静的ファイルとして配信します。"
            ]
        },
        {
            "title": "機能紹介④：Mikka AI Assistant（Gemini）",
            "img": "screenshot_chat.png",
            "bullets": [
                "【特徴】データベースと連携したスマートな飲食店推薦AI",
                "・Google Gemini APIのシームレスな連携",
                "  右下のフローティングAIチャットウィジェットからいつでも利用可能。",
                "・RAG（検索拡張生成）ライクなコンテキスト注入",
                "  DB内の最新店舗情報をプロンプトに動的に注入することで、",
                "  実在するおすすめ店舗を自然言語で提案させます（ハルシネーションの防止）。",
                "・便利なクイックリプライ（提案ピル）",
                "  「Milliy taomlar 🇺🇿」等のボタンでワンタップで質問可能です。"
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # 背景色の設定
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(245, 247, 250)
        
        # スライドタイトルの追加
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.name = 'MS Gothic'
        p.font.color.rgb = RGBColor(30, 41, 59)
        
        # 左側にスクリーンショット画像を配置
        img_path = os.path.join(img_dir, data["img"])
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(0.5), Inches(1.2), width=Inches(4.5), height=Inches(2.8))
            
        # 右側に詳しい説明テキストを配置
        descBox = slide.shapes.add_textbox(Inches(5.2), Inches(1.2), Inches(4.3), Inches(5.5))
        desc_tf = descBox.text_frame
        desc_tf.word_wrap = True
        
        for i, bullet in enumerate(data["bullets"]):
            if i == 0:
                p_bullet = desc_tf.paragraphs[0]
                p_bullet.text = bullet
                p_bullet.font.size = Pt(14)
                p_bullet.font.bold = True
                p_bullet.font.color.rgb = RGBColor(30, 41, 59)
            else:
                p_bullet = desc_tf.add_paragraph()
                p_bullet.text = bullet
                if bullet.strip().startswith("・") or bullet.strip().startswith("【"):
                    p_bullet.font.size = Pt(12)
                    p_bullet.font.bold = True
                    p_bullet.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    p_bullet.font.size = Pt(11)
                    p_bullet.font.color.rgb = RGBColor(71, 85, 105)
            p_bullet.font.name = 'MS Gothic'
            p_bullet.space_after = Pt(4)
            
    prs.save(ppt_path)
    print("Successfully inserted 4 detailed slides with explanations into Mikka_Presentation.pptx.")

if __name__ == "__main__":
    main()
