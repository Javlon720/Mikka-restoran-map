from pptx import Presentation
import os

ppt_path = "/Users/macbook-projavlon/Desktop/JDU_bitiruv/employability/abdullayevJavlon214773.pptx"

def main():
    if not os.path.exists(ppt_path):
        print(f"{ppt_path} not found.")
        return

    prs = Presentation(ppt_path)
    
    # Slide 4 (Index 3) - Strengths 2 (Design & Japanese)
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            for p in tf.paragraphs:
                # Replace N3 with N2
                if "日本語N3" in p.text:
                    p.text = p.text.replace("日本語N3", "日本語N2 (JLPT N2合格)")
                    p.font.name = 'MS Gothic'
                    print("Updated Slide 4 bullet point for N2 qualification.")
                elif "N3レベル" in p.text:
                    p.text = p.text.replace("N3レベル", "N2合格")
                    p.font.name = 'MS Gothic'

    # Update Slide 4 speaker notes
    notes_slide4 = slide4.notes_slide
    notes_tf = notes_slide4.notes_text_frame
    original_notes = notes_tf.text
    
    # Replace N3/N2 study statements with N2/N1 study
    updated_notes = original_notes.replace(
        "にほんごは エヌさん（N3）レベルで、げんざい エヌに（N2）の べんきょうを すすめています。",
        "にほんごは エヌに（N2）レベルに ごうかくしており、げんざい エヌいち（N1）の べんきょうを すすめています。"
    )
    notes_tf.text = updated_notes
    print("Updated Slide 4 speaker notes for N2 level.")
    
    # Check Slide 5 notes just in case
    slide5 = prs.slides[4]
    notes_slide5 = slide5.notes_slide
    notes_tf5 = notes_slide5.notes_text_frame
    original_notes5 = notes_tf5.text
    updated_notes5 = original_notes5.replace("エヌさん", "エヌに").replace("N3", "N2")
    notes_tf5.text = updated_notes5

    prs.save(ppt_path)
    print("Successfully updated language level in abdullayevJavlon214773.pptx!")

if __name__ == "__main__":
    main()
